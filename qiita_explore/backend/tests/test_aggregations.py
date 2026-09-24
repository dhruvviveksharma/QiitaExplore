"""Sample Aggregation: SQLite CRUD (store/aggregation_crud) and the
/api/aggregations routes with Qiita Postgres stubbed. The pure CSV row builder
is covered in test_fastq_manifest.py."""

import csv
import os
import sys

import pytest

from .conftest import stub_qiita_db_and_core


# ── store ────────────────────────────────────────────────────────────────────

@pytest.fixture
def agg_crud():
    import store.aggregation_crud as m
    return m


STUDY = {"study_id": 16326, "study_title": "Test", "data_types": "16S", "num_samples": 2, "num_preps": 1,
         "study_abstract": "About soil", "pi_name": "Rob Knight", "pi_affiliation": "UCSD",
         "year": 2015, "is_gold": True}
SAMPLES = ["s1", "s2"]


def test_store_create_list_get(agg_crud):
    agg = agg_crud.create_aggregation("u1", "A")
    listed = agg_crud.list_aggregations("u1")
    assert [a["aggregation_id"] for a in listed] == [agg["aggregation_id"]]
    assert listed[0]["studies"] == []
    assert agg_crud.get_aggregation(agg["aggregation_id"], "u2") is None
    assert agg_crud.list_aggregations("u2") == []


def test_store_add_readd_remove_study(agg_crud):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    agg = agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 2, SAMPLES)
    assert len(agg["studies"]) == 1
    s = agg["studies"][0]
    assert (s["study_id"], s["study_title"], s["data_types"], s["num_samples"], s["num_preps"],
            s["fastq_artifact_count"]) == (16326, "Test", "16S", 2, 1, 2)
    # header snapshot for the tab's cards
    assert (s["study_abstract"], s["pi_name"], s["pi_affiliation"], s["year"], s["is_gold"]) == \
        ("About soil", "Rob Knight", "UCSD", 2015, 1)
    assert s["selected_samples"] == 2
    assert len(agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 2, SAMPLES)["studies"]) == 1
    assert agg_crud.add_study_to_aggregation(aid, "u2", STUDY, 2, SAMPLES) is None
    assert agg_crud.remove_study_from_aggregation(aid, "u2", 16326) is None
    assert agg_crud.remove_study_from_aggregation(aid, "u1", 16326)["studies"] == []


def test_store_set_samples_add_remove_clear(agg_crud):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 2, SAMPLES)

    agg = agg_crud.set_aggregation_samples(aid, "u1", 16326, remove=["s2"])
    assert agg["studies"][0]["selected_samples"] == 1
    assert agg_crud.selected_in(aid, 16326, ["s1", "s2", "zz"]) == {"s1"}

    agg = agg_crud.set_aggregation_samples(aid, "u1", 16326, add=["s2", "s3"])
    assert agg["studies"][0]["selected_samples"] == 3
    assert agg_crud.selected_by_study(aid) == {16326: {"s1", "s2", "s3"}}

    agg = agg_crud.set_aggregation_samples(aid, "u1", 16326, clear=True, add=["s9"])
    assert agg["studies"][0]["selected_samples"] == 1
    assert agg_crud.selected_in(aid, 16326, ["s9"]) == {"s9"}
    assert agg_crud.selected_in(aid, 16326, []) == set()

    assert agg_crud.set_aggregation_samples(aid, "u1", 16326, clear=True)["studies"][0]["selected_samples"] == 0
    assert agg_crud.set_aggregation_samples(aid, "u2", 16326, add=["s1"]) is None


def test_store_readd_keeps_deselection(agg_crud):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 2, SAMPLES)
    agg_crud.set_aggregation_samples(aid, "u1", 16326, remove=["s2"])
    # Re-adding a study already present must not re-check what the user unchecked.
    agg = agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 2, SAMPLES)
    assert agg["studies"][0]["selected_samples"] == 1


def test_store_rename(agg_crud):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    assert agg_crud.rename_aggregation(aid, "u1", "B")["name"] == "B"
    assert agg_crud.rename_aggregation(aid, "u2", "C") is None
    assert agg_crud.get_aggregation(aid, "u1")["name"] == "B"


def test_store_delete_cascades(agg_crud, db_conn):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 1, SAMPLES)
    assert db_conn.execute("SELECT COUNT(*) FROM aggregation_samples").fetchone()[0] == 2
    assert agg_crud.delete_aggregation(aid, "u2") is False
    assert agg_crud.delete_aggregation(aid, "u1") is True
    assert agg_crud.list_aggregations("u1") == []
    assert db_conn.execute("SELECT COUNT(*) FROM aggregation_studies").fetchone()[0] == 0
    # chain cascade: aggregation -> studies -> samples
    assert db_conn.execute("SELECT COUNT(*) FROM aggregation_samples").fetchone()[0] == 0


def test_store_remove_study_cascades_samples(agg_crud, db_conn):
    aid = agg_crud.create_aggregation("u1", "A")["aggregation_id"]
    agg_crud.add_study_to_aggregation(aid, "u1", STUDY, 1, SAMPLES)
    agg_crud.remove_study_from_aggregation(aid, "u1", 16326)
    assert db_conn.execute("SELECT COUNT(*) FROM aggregation_samples").fetchone()[0] == 0


# ── routes ───────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def _app(tmp_path_factory):
    db_path = str(tmp_path_factory.mktemp("aggregations") / "test.db")
    os.environ["QIITA_EXPERIMENT_DB_PATH"] = db_path
    for name in list(sys.modules):
        if (name == "run" or name.startswith("routes.") or name == "store"
                or name.startswith("store.") or name.startswith("helpers.")
                or "sql_store" in name):
            del sys.modules[name]
    stub_qiita_db_and_core()
    import run
    import config
    config.ALLOWED_ORIGINS = []
    config.SESSION_COOKIE_SECURE = False
    return run.app


@pytest.fixture
def client(_app):
    return _app.test_client()


@pytest.fixture
def logged_in(client, monkeypatch):
    """Session cookie on `client`; returns the CSRF header dict for state-changing calls."""
    import routes.auth_routes as auth_routes
    from helpers.qiita_client import WhoAmIResult

    monkeypatch.setattr(auth_routes, "whoami", lambda pat: WhoAmIResult(ok=True, identity={
        "principal_idx": 90003, "email": "agg@test.local",
        "system_role": "user", "scopes": [], "profile_complete": True,
    }))
    resp = client.post("/api/auth/connect", json={"token": "qk_test"})
    assert resp.status_code == 200, resp.get_json()
    return {"X-CSRF-Token": resp.get_json()["csrf_token"]}


@pytest.fixture
def stub_qiita(monkeypatch):
    """Every Postgres-touching name the routes import, patched on the route
    module. s2 is the one sample with a file (fastq paired), so tests can
    exercise files-first ordering / show filtering / "select: with_files"
    without a real availability computation."""
    import routes.aggregation_routes as ar
    _FIELDS = {"s1": ("s1", "stool"), "s2": ("s2", "skin")}
    monkeypatch.setattr(ar, "is_study_public", lambda sid: sid != 99999)
    monkeypatch.setattr(ar, "count_fastq_artifacts", lambda sid: 2)
    monkeypatch.setattr(ar, "list_study_sample_ids", lambda sid: ["s1", "s2"])
    monkeypatch.setattr(ar, "matching_sample_ids", lambda sid, q: ["s2"])
    monkeypatch.setattr(ar, "display_columns", lambda sid: ["sample_type"])
    monkeypatch.setattr(ar, "fetch_samples_by_ids", lambda sid, ids: [_FIELDS[i] for i in ids if i in _FIELDS])
    monkeypatch.setattr(ar, "get_sample_files", lambda sid: {"s2": [("16S", "Raw upload", 2, 0)]})
    monkeypatch.setattr(ar, "fetch_aggregate_csv_rows", lambda selected, file_filter=None: [
        (16326, "s1", "/a/f_R1.fq.gz", "16S", "raw_forward_seqs", "Raw upload"),
        (16326, "s1", "/a/f_R2.fq.gz", "16S", "raw_reverse_seqs", "Raw upload"),
    ])
    return ar


def _create(client, hdr, name="Agg"):
    r = client.post("/api/aggregations", json={"name": name}, headers=hdr)
    assert r.status_code == 201, r.get_json()
    return r.get_json()


def _add(client, hdr, aid, study_id=16326):
    return client.post(f"/api/aggregations/{aid}/studies", headers=hdr,
                       json={"study": {**STUDY, "study_id": study_id}})


def _listed_ids(client):
    return [a["aggregation_id"] for a in client.get("/api/aggregations").get_json()["aggregations"]]


def test_route_crud_roundtrip(client, logged_in, stub_qiita):
    agg = _create(client, logged_in, "Round trip")
    aid = agg["aggregation_id"]
    assert agg["studies"] == []
    assert aid in _listed_ids(client)

    r = client.patch(f"/api/aggregations/{aid}", json={"name": "Renamed"}, headers=logged_in)
    assert r.status_code == 200 and r.get_json()["name"] == "Renamed"

    r = _add(client, logged_in, aid)
    assert r.status_code == 200, r.get_json()
    studies = r.get_json()["studies"]
    assert [s["study_id"] for s in studies] == [16326]
    assert studies[0]["fastq_artifact_count"] == 2
    # whole study added = every sample checked; badge denominator = rows stored
    assert (studies[0]["selected_samples"], studies[0]["num_samples"]) == (2, 2)
    assert studies[0]["pi_name"] == "Rob Knight"

    r = client.delete(f"/api/aggregations/{aid}/studies/16326", headers=logged_in)
    assert r.status_code == 200 and r.get_json()["studies"] == []

    r = client.delete(f"/api/aggregations/{aid}", headers=logged_in)
    assert r.status_code == 200 and r.get_json() == {"deleted": aid}
    assert aid not in _listed_ids(client)


def test_route_rename_blank_400_and_unknown_404(client, logged_in, stub_qiita):
    aid = _create(client, logged_in)["aggregation_id"]
    assert client.patch(f"/api/aggregations/{aid}", json={"name": "  "}, headers=logged_in).status_code == 400
    assert client.patch("/api/aggregations/nope", json={"name": "x"}, headers=logged_in).status_code == 404
    assert client.delete("/api/aggregations/nope", headers=logged_in).status_code == 404


def test_route_add_study_errors(client, logged_in, stub_qiita, monkeypatch):
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid, study_id=99999).status_code == 403
    assert _add(client, logged_in, "nope").status_code == 404
    r = client.post(f"/api/aggregations/{aid}/studies", json={"study": {}}, headers=logged_in)
    assert r.status_code == 400
    monkeypatch.setattr(stub_qiita, "AGGREGATION_STUDIES_CAP", 1)
    assert _add(client, logged_in, aid, study_id=1).status_code == 200
    r = _add(client, logged_in, aid, study_id=2)
    assert r.status_code == 400 and "maximum" in r.get_json()["error"]


def test_route_samples_page(client, logged_in, stub_qiita):
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    client.patch(f"/api/aggregations/{aid}/studies/16326/samples", json={"remove": ["s2"]}, headers=logged_in)

    r = client.get(f"/api/aggregations/{aid}/studies/16326/samples?offset=0&limit=50")
    assert r.status_code == 200, r.get_json()
    page = r.get_json()
    assert (page["study_id"], page["total"], page["offset"], page["limit"]) == (16326, 2, 0, 50)
    assert page["columns"] == ["sample_type"]
    assert page["selected_count"] == 1
    assert page["with_files"] == 1
    # files-first: s2 (has a file) sorts before s1, id order preserved within each group
    assert [row["sample_id"] for row in page["rows"]] == ["s2", "s1"]
    assert page["rows"][0] == {"sample_id": "s2", "selected": False, "fastq": "paired", "fasta": False,
                                "data_types": ["16S"], "processing": ["Raw upload"],
                                "fields": {"sample_type": "skin"}}
    assert page["rows"][1] == {"sample_id": "s1", "selected": True, "fastq": None, "fasta": False,
                                "data_types": [], "processing": [], "fields": {"sample_type": "stool"}}

    only_files = client.get(f"/api/aggregations/{aid}/studies/16326/samples?show=with_files").get_json()
    assert [row["sample_id"] for row in only_files["rows"]] == ["s2"]
    assert only_files["total"] == 1

    without_files = client.get(f"/api/aggregations/{aid}/studies/16326/samples?show=without_files").get_json()
    assert [row["sample_id"] for row in without_files["rows"]] == ["s1"]
    assert without_files["total"] == 1

    assert client.get(f"/api/aggregations/{aid}/studies/16326/samples?show=bogus").status_code == 400

    # q routes to matching_sample_ids (stub returns only s2, regardless of q)
    q_page = client.get(f"/api/aggregations/{aid}/studies/16326/samples?q=x").get_json()
    assert [row["sample_id"] for row in q_page["rows"]] == ["s2"]
    assert q_page["total"] == 1

    # limit is clamped, offset floors at 0
    page = client.get(f"/api/aggregations/{aid}/studies/16326/samples?limit=9999&offset=-5").get_json()
    assert (page["limit"], page["offset"]) == (500, 0)

    assert client.get(f"/api/aggregations/{aid}/studies/16326/samples?offset=x").status_code == 400
    assert client.get(f"/api/aggregations/{aid}/studies/777/samples").status_code == 404
    assert client.get("/api/aggregations/nope/studies/16326/samples").status_code == 404


def test_route_set_samples_variants(client, logged_in, stub_qiita):
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    url = f"/api/aggregations/{aid}/studies/16326/samples"

    def count(resp):
        assert resp.status_code == 200, resp.get_json()
        return resp.get_json()["studies"][0]["selected_samples"]

    assert count(client.patch(url, json={"select": "none"}, headers=logged_in)) == 0
    assert count(client.patch(url, json={"add": ["s1"]}, headers=logged_in)) == 1
    assert count(client.patch(url, json={"select": "matching", "q": "skin"}, headers=logged_in)) == 2
    assert count(client.patch(url, json={"remove": ["s1", "s2"]}, headers=logged_in)) == 0
    assert count(client.patch(url, json={"select": "all"}, headers=logged_in)) == 2
    # "with_files" replaces the whole selection with only the samples that
    # resolve to a file — exactly what the CSV export can contain.
    assert count(client.patch(url, json={"select": "with_files"}, headers=logged_in)) == 1

    for bad in [{}, {"add": "s1"}, {"add": [1]}, {"select": "matching"}, {"select": "some"}]:
        r = client.patch(url, json=bad, headers=logged_in)
        assert r.status_code == 400, bad
    assert client.patch(f"/api/aggregations/{aid}/studies/777/samples",
                        json={"select": "all"}, headers=logged_in).status_code == 404
    assert client.patch(url, json={"select": "all"}).status_code == 403   # no CSRF header


def test_route_export_csv(client, logged_in, stub_qiita):
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    resp = client.get(f"/api/aggregations/{aid}/export.csv")
    assert resp.status_code == 200
    assert resp.mimetype == "text/csv"
    assert resp.headers["Content-Disposition"] == f"attachment; filename=aggregation_{aid}_samples.csv"
    rows = list(csv.reader(resp.get_data(as_text=True).splitlines()))
    assert rows[0] == ["study_id", "sample_id", "file_path_in_qmounts", "data_type", "file_type", "processing"]
    assert rows[1] == ["16326", "s1", "/a/f_R1.fq.gz", "16S", "raw_forward_seqs", "Raw upload"]
    assert rows[2] == ["16326", "s1", "/a/f_R2.fq.gz", "16S", "raw_reverse_seqs", "Raw upload"]


def test_route_export_xlsx(client, logged_in, stub_qiita):
    import io
    import openpyxl
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    resp = client.get(f"/api/aggregations/{aid}/export.xlsx")
    assert resp.status_code == 200, resp.get_data(as_text=True)[:200]
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.headers["Content-Disposition"] == f"attachment; filename=aggregation_{aid}_samples.xlsx"
    ws = openpyxl.load_workbook(io.BytesIO(resp.get_data())).active
    assert [c.value for c in ws[2]] == [16326, "s1", "/a/f_R1.fq.gz", "16S", "raw_forward_seqs", "Raw upload"]
    assert ws.cell(row=2, column=2).data_type == "s"
    assert client.get("/api/aggregations/nope/export.xlsx").status_code == 404


def test_route_export_passes_only_checked_samples(client, logged_in, stub_qiita, monkeypatch):
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    client.patch(f"/api/aggregations/{aid}/studies/16326/samples", json={"remove": ["s2"]}, headers=logged_in)
    seen = {}

    def _capture(selected, file_filter=None):
        seen.update(selected)
        seen["filter"] = file_filter
        return [(16326, "s1", "/a/f.fq.gz", "16S", "raw_forward_seqs", "Raw upload")]
    monkeypatch.setattr(stub_qiita, "fetch_aggregate_csv_rows", _capture)
    filt = {"data_types": ["16S"], "processing": ["Raw upload"]}
    client.patch(f"/api/aggregations/{aid}", json={"file_filter": filt}, headers=logged_in)
    assert client.get(f"/api/aggregations/{aid}/export.csv").status_code == 200
    assert seen == {16326: {"s1"}, "filter": filt}


def test_route_export_errors(client, logged_in, stub_qiita, monkeypatch):
    aid = _create(client, logged_in)["aggregation_id"]
    assert client.get(f"/api/aggregations/{aid}/export.csv").status_code == 400   # no studies
    assert client.get("/api/aggregations/nope/export.csv").status_code == 404
    assert _add(client, logged_in, aid).status_code == 200
    client.patch(f"/api/aggregations/{aid}/studies/16326/samples", json={"select": "none"}, headers=logged_in)
    r = client.get(f"/api/aggregations/{aid}/export.csv")
    assert r.status_code == 400 and r.get_json()["error"] == "No samples selected"

    def _raise(selected, file_filter=None):
        raise ValueError("None of the selected samples has a per-sample sequence file")
    monkeypatch.setattr(stub_qiita, "fetch_aggregate_csv_rows", _raise)
    client.patch(f"/api/aggregations/{aid}/studies/16326/samples", json={"select": "all"}, headers=logged_in)
    resp = client.get(f"/api/aggregations/{aid}/export.csv")
    assert resp.status_code == 404
    assert resp.get_json()["error"] == "None of the selected samples has a per-sample sequence file"


# ── the saved file filter ────────────────────────────────────────────────────

_FILTER_MAP = {"s1": [("Metagenomic", "Atropos v1.1.24", 2, 0)], "s2": [("16S", "Raw upload", 0, 1)]}


def test_route_file_filter_patch_validates_and_persists(client, logged_in, stub_qiita):
    aid = _create(client, logged_in)["aggregation_id"]
    assert client.get("/api/aggregations").get_json()["aggregations"][0]["file_filter"] == \
        {"data_types": [], "processing": []}
    url = f"/api/aggregations/{aid}"
    r = client.patch(url, json={"file_filter": {"data_types": ["16S", "16S"]}}, headers=logged_in)
    assert r.status_code == 200, r.get_json()
    assert r.get_json()["file_filter"] == {"data_types": ["16S"], "processing": []}
    listed = [a for a in client.get("/api/aggregations").get_json()["aggregations"] if a["aggregation_id"] == aid]
    assert listed[0]["file_filter"] == {"data_types": ["16S"], "processing": []}
    r = client.patch(url, json={"name": "Both", "file_filter": {}}, headers=logged_in)
    assert (r.get_json()["name"], r.get_json()["file_filter"]) == ("Both", {"data_types": [], "processing": []})
    for bad in [{}, {"file_filter": []}, {"file_filter": {"other": []}}, {"file_filter": {"data_types": "16S"}},
                {"file_filter": {"data_types": [""]}}, {"file_filter": {"processing": ["x" * 201]}},
                {"file_filter": {"data_types": [str(i) for i in range(51)]}}]:
        assert client.patch(url, json=bad, headers=logged_in).status_code == 400, bad
    assert client.patch("/api/aggregations/nope", json={"file_filter": {}}, headers=logged_in).status_code == 404


def test_route_samples_and_select_follow_file_filter(client, logged_in, stub_qiita, monkeypatch):
    monkeypatch.setattr(stub_qiita, "get_sample_files", lambda sid: _FILTER_MAP)
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    base = f"/api/aggregations/{aid}/studies/16326/samples"

    page = client.get(base).get_json()               # no filter: both have files
    assert page["with_files"] == 2
    client.patch(f"/api/aggregations/{aid}", json={"file_filter": {"data_types": ["16S"]}}, headers=logged_in)
    page = client.get(base).get_json()
    assert page["with_files"] == 1
    assert [r["sample_id"] for r in page["rows"]] == ["s2", "s1"]  # files-first under the filter
    s2, s1 = page["rows"]
    assert (s2["fastq"], s2["fasta"]) == (None, True)
    # s1's metagenomic file is filtered out, but still reported so the UI can dim it
    assert (s1["fastq"], s1["fasta"], s1["data_types"], s1["processing"]) == \
        (None, False, ["Metagenomic"], ["Atropos v1.1.24"])
    assert [r["sample_id"] for r in client.get(base + "?show=without_files").get_json()["rows"]] == ["s1"]

    r = client.patch(base, json={"select": "with_files"}, headers=logged_in)
    assert r.get_json()["studies"][0]["selected_samples"] == 1


def test_route_file_facets(client, logged_in, stub_qiita, monkeypatch):
    monkeypatch.setattr(stub_qiita, "get_sample_files", lambda sid: _FILTER_MAP)
    aid = _create(client, logged_in)["aggregation_id"]
    assert _add(client, logged_in, aid).status_code == 200
    d = client.get(f"/api/aggregations/{aid}/file-facets").get_json()
    assert d["data_types"] == [{"name": "16S", "count": 1}, {"name": "Metagenomic", "count": 1}]
    assert d["processing"] == [{"name": "Atropos v1.1.24", "count": 1}, {"name": "Raw upload", "count": 1}]
    assert (d["exportable"], d["selected"]) == (2, 2)

    client.patch(f"/api/aggregations/{aid}", json={"file_filter": {"data_types": ["Metagenomic"]}},
                 headers=logged_in)
    client.patch(f"/api/aggregations/{aid}/studies/16326/samples", json={"remove": ["s1"]}, headers=logged_in)
    d = client.get(f"/api/aggregations/{aid}/file-facets").get_json()
    assert d["processing"] == [{"name": "Atropos v1.1.24", "count": 1}]  # narrowed by the data-type pick
    assert (d["exportable"], d["selected"]) == (0, 1)  # s2 is checked but has no metagenomic file
    assert client.get("/api/aggregations/nope/file-facets").status_code == 404


def test_route_401_without_session(_app):
    assert _app.test_client().get("/api/aggregations").status_code == 401


def test_route_post_requires_csrf(client, logged_in):
    assert client.post("/api/aggregations", json={"name": "x"}).status_code == 403


# ── TKT-086: a study_detail_cache row without preps is a miss ────────────────

@pytest.mark.parametrize("sid, preps_json", [(70001, None), (70002, "[]")])
def test_study_detail_refetches_when_cached_row_has_no_preps(client, logged_in, monkeypatch, sid, preps_json):
    """Other writers create study_detail_cache rows without preps, and the old
    availability map persisted "[]" into some; the modal must refetch rather
    than serve zero preps as a hit."""
    import routes.study_routes as sr

    # Write through the route module's own binding: fresh_db re-imports
    # store per test, but this module-scoped app still holds the original.
    sr.upsert_study_detail_cache(sid, preps_json, preps_json, samples_context="ctx")
    fetched = []
    monkeypatch.setattr(sr, "is_study_public", lambda sid: True)
    monkeypatch.setattr(sr, "_fetch_study_detail_from_qiita",
                        lambda sid: (fetched.append(sid) or ([{"prep_template_id": 7, "data_type": "16S"}], [])))
    monkeypatch.setattr(sr, "fetch_artifact_graph", lambda sid: [])
    monkeypatch.setattr(sr, "_fetch_prep_metadata_summary", lambda pid: {})
    monkeypatch.setattr(sr, "_fetch_study_samples", lambda sid, limit=200: ([], 0))
    monkeypatch.setattr(sr, "_fetch_sample_context_text", lambda sid: "")
    r = client.get(f"/api/studies/{sid}/detail")
    assert r.status_code == 200, r.get_json()
    assert fetched == [sid]
    assert [p["prep_template_id"] for p in r.get_json()["preps"]] == [7]
